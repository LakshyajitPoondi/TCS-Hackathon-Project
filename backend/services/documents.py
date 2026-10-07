"""Document extraction, scoped mapping and provenance-preserving chunks."""
import io
import re
import uuid
from pathlib import Path
from fastapi import HTTPException
from sqlalchemy import select
from backend import db
from backend.core import config

DOC_TYPES={'manual','sop','maintenance_guide','troubleshooting_guide','spec_sheet','other'}
SCOPES={'machine','model','line','plant'}

def extract(filename,content):
    suffix=Path(filename).suffix.lower()
    if suffix not in ('.pdf','.docx','.md','.txt'): raise HTTPException(422,'Supported documents: PDF, DOCX, Markdown, TXT')
    if len(content)>config.DOC_UPLOAD_MAX_MB*1024*1024: raise HTTPException(413,'Document exceeds DOC_UPLOAD_MAX_MB')
    try:
        if suffix=='.pdf':
            from pypdf import PdfReader
            parts=[{'location':f'Page {i}', 'text':page.extract_text() or ''} for i,page in enumerate(PdfReader(io.BytesIO(content)).pages,1)]
        elif suffix=='.docx':
            from docx import Document
            parts=[]; heading='Document'; lines=[]
            for p in Document(io.BytesIO(content)).paragraphs:
                if p.style.name.startswith('Heading'):
                    if lines: parts.append({'location':heading,'text':'\n'.join(lines)})
                    heading=p.text; lines=[]
                else: lines.append(p.text)
            if lines: parts.append({'location':heading,'text':'\n'.join(lines)})
        else:
            text=content.decode('utf-8-sig');parts=[];heading='Document';lines=[]
            for line in text.splitlines():
                if re.match(r'^#{1,6}\s',line):
                    if lines:parts.append({'location':heading,'text':'\n'.join(lines)})
                    heading=line.lstrip('# ').strip();lines=[]
                else: lines.append(line)
            if lines:parts.append({'location':heading,'text':'\n'.join(lines)})
    except Exception as exc:
        raise HTTPException(422,f'Cannot extract document: {type(exc).__name__}') from exc
    parts=[p for p in parts if p['text'].strip()]
    if not parts:raise HTTPException(422,'Document has no extractable text (OCR is not supported)')
    return parts

def detect(filename,parts,machines):
    suggested=[];ambiguous=[]
    locations=[{'location':'Filename','text':filename}]+parts
    for part in locations:
        text=part['text'];explicit=set()
        for m in machines:
            pattern=re.escape(m.line)+r'[/_\s-]+'+re.escape(m.short_name)
            if re.search(pattern,text,re.I) or (m.data.get('serial_number') and m.data['serial_number'].lower() in text.lower()):
                explicit.add(m.machine_uid)
                suggested.append({'scope':'machine','target':m.machine_uid,'matched_terms':[m.machine_uid], 'location':part['location'],'confidence':'high'})
        for short in sorted({m.short_name for m in machines}):
            if re.search(r'(?<![\w-])'+re.escape(short)+r'(?![\w-])',text,re.I) and not any(m.short_name==short and m.machine_uid in explicit for m in machines):
                candidates=[m.machine_uid for m in machines if m.short_name==short]
                ambiguous.append({'short_name':short,'candidates':candidates,'model_scopes':sorted({m.model for m in machines if m.short_name==short}),'location':part['location']})
        for line in sorted({m.line for m in machines}):
            if line.lower() in text.lower(): suggested.append({'scope':'line','target':line,'matched_terms':[line],'location':part['location'],'confidence':'medium'})
        for model in sorted({m.model for m in machines}):
            if re.search(r'\b'+re.escape(model)+r'\b',text,re.I): suggested.append({'scope':'model','target':model,'matched_terms':[model],'location':part['location'],'confidence':'medium'})
        for maker in sorted({m.data.get('manufacturer','') for m in machines}-{'' ,'Unknown (synthetic)'}):
            if maker.lower() in text.lower(): suggested.append({'scope':'manufacturer','target':maker,'matched_terms':[maker],'location':part['location'],'confidence':'low'})
    return {'suggestions':suggested,'ambiguities':ambiguous,'ambiguous':bool(ambiguous),'warnings':[],'conflict':False}

def machine_allowed(machine,scope,targets):
    return scope=='plant' or (scope=='machine' and machine.machine_uid in targets) or (scope=='line' and machine.line in targets) or (scope=='model' and machine.model in targets)

def validate_mapping(scope,targets,machines):
    if scope not in SCOPES:raise HTTPException(422,'Invalid document scope')
    valid={'machine':{m.machine_uid for m in machines},'line':{m.line for m in machines},'model':{m.model for m in machines},'plant':set()}[scope]
    if scope!='plant' and (not targets or not set(targets)<=valid):raise HTTPException(422,'Choose valid scope targets')
    if scope=='plant' and targets:raise HTTPException(422,'Plant scope has no targets')

def put_links(session,doc,machines):
    from sqlalchemy import delete
    session.execute(delete(db.DocumentMachineLink).where(db.DocumentMachineLink.doc_id==doc.doc_id))
    for m in machines:
        if machine_allowed(m,doc.scope,doc.targets):session.add(db.DocumentMachineLink(doc_id=doc.doc_id,machine_uid=m.machine_uid))

def make_chunks(doc_id,version,parts):
    chunks=[]
    for part in parts:
        words=part['text'].split()
        for start in range(0,len(words),340):
            text=' '.join(words[start:start+400])
            if text: chunks.append(db.DocumentChunk(chunk_id=f'{doc_id}:v{version}:{len(chunks)+1}',doc_id=doc_id,version=version,page_or_section=part['location'],text=text))
    return chunks

def ingest(session,filename,content,title,doc_type,version,scope,targets,user_id,doc_id=None):
    if doc_type not in DOC_TYPES: raise HTTPException(422,'Invalid doc_type')
    if not title.strip() or len(title)>300 or not version or len(version)>50:raise HTTPException(422,'Invalid title/version')
    machines=list(session.scalars(select(db.Machine)))
    validate_mapping(scope,targets,machines)
    parts=extract(filename,content);detection=detect(filename,parts,machines)
    detected={s['target'] for s in detection['suggestions'] if s['scope']=='machine'}
    allowed={m.machine_uid for m in machines if machine_allowed(m,scope,targets)}
    if detected-allowed:
        detection['conflict']=True;detection['warnings'].append('Explicit mapping disagrees with detected machine identity')
    # Filename and content machine suggestions must agree even for broad scopes.
    names={s['target'] for s in detection['suggestions'] if s['scope']=='machine' and s['location']=='Filename'}
    bodies={s['target'] for s in detection['suggestions'] if s['scope']=='machine' and s['location']!='Filename'}
    if names and bodies and names!=bodies:
        detection['conflict']=True;detection['warnings'].append('Filename and document text identify different machines')
    if detection['ambiguous']:detection['warnings'].append('Short machine name without line requires confirmation')
    doc_id=doc_id or 'DOC-'+uuid.uuid4().hex
    if session.get(db.Document,doc_id):raise HTTPException(409,'Document ID already exists')
    config.DOCUMENTS_DIR.mkdir(parents=True,exist_ok=True)
    path=config.DOCUMENTS_DIR/(uuid.uuid4().hex+Path(filename).suffix.lower());path.write_bytes(content)
    doc=db.Document(doc_id=doc_id,title=title.strip(),doc_type=doc_type,version=version,scope=scope,targets=targets,status='pending_mapping' if detection['ambiguous'] or detection['conflict'] else 'active',uploaded_by=user_id,detection=detection,storage_path=str(path))
    session.add(doc);session.flush();session.add_all(make_chunks(doc_id,version,parts));put_links(session,doc,machines)
    return doc

def serialize(doc):
    return {name:getattr(doc,name) for name in ('doc_id','title','doc_type','version','scope','targets','status','uploaded_by','uploaded_at','detection')}
