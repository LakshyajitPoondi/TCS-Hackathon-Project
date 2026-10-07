import io
from backend import db
from sqlalchemy import select,func
from backend.services.documents import extract

def test_registry_seed(client):
    rows=client.get('/api/machines').json()
    assert len(rows)==9
    assert rows[0]['machine_uid']=='LINE-A/IMM-01'
    assert 'normal_ranges' in rows[0]
    assert len(client.get('/api/machines/LINE-A/IMM-01/documents').json())==8
    assert 0<len(client.get('/api/machines/LINE-A/IMM-01/incidents').json())<6   # B12: affected machine only
    assert client.get('/api/machines/LINE-A/IMM-01').status_code==200

def upload(client,text,scope='machine',targets='["LINE-A/IMM-01"]',filename='guide.md'):
    return client.post('/api/documents/upload',data={'title':'Test guide','doc_type':'manual','scope':scope,'targets':targets},files={'file':(filename,text.encode())})

def test_mapping(client):
    r=upload(client,'# Maintenance\nIMM-01 cooling checks.')
    assert r.status_code==201
    d=r.json();assert d['status']=='pending_mapping'
    assert len(d['detection']['ambiguities'][0]['candidates'])==3
    assert client.post('/api/documents/'+d['doc_id']+'/confirm',json={'scope':'machine','targets':['LINE-A/IMM-01']}).json()['status']=='active'
    r=upload(client,'# Manual\nLINE-B/IMM-01 cooling checks.',filename='LINE-A_IMM-01.md')
    assert r.json()['detection']['conflict'] and r.json()['status']=='pending_mapping'
    r=upload(client,'# Manual\nLINE-A/IMM-01 cooling checks.')
    assert r.json()['status']=='active'
    detail=client.get('/api/documents/'+r.json()['doc_id']).json()
    assert detail['chunks'][0]['page_or_section']=='Manual'
    assert client.get('/api/documents?machine_uid=LINE-A/IMM-01&doc_type=manual&status=active').status_code==200

def test_docx_pdf_and_limits(client,monkeypatch):
    from docx import Document
    doc=Document();doc.add_heading('Cooling',1);doc.add_paragraph('Check flow.');buffer=io.BytesIO();doc.save(buffer)
    assert extract('manual.docx',buffer.getvalue())[0]['location']=='Cooling'
    from pypdf import PdfWriter
    writer=PdfWriter();writer.add_blank_page(width=100,height=100);buffer=io.BytesIO();writer.write(buffer)
    from fastapi import HTTPException
    import pytest
    with pytest.raises(HTTPException):extract('blank.pdf',buffer.getvalue())
    from backend.core import config
    monkeypatch.setattr(config,'DOC_UPLOAD_MAX_MB',.000001)
    assert upload(client,'too large').status_code==413

def test_doc_permissions(client,tokens):
    for role in ('qa_lead','viewer'):
        client.headers['Authorization']='Bearer '+tokens[role]
        assert upload(client,'text').status_code==403
        assert client.post('/api/documents/SOP-007/confirm',json={'scope':'plant','targets':[]}).status_code==403
